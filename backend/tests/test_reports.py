"""Tests for the reporting & export module: generation of every report type in
every format, persistence, file handling, downloads, RBAC and ownership,
validation, cohort guards, deletion and retention purge.
"""

import csv
import io
from datetime import UTC, date, datetime, timedelta
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core import rbac
from app.core.config import settings
from app.core.rbac import Permission, Role
from app.models.admission import Admission
from app.models.audit_log import AuditLog
from app.models.patient import Patient
from app.models.prediction import RiskPrediction
from app.models.report import REPORT_FORMATS, REPORT_TYPES, Report
from app.models.treatment import TreatmentOutcome
from app.utils import report_storage
from app.utils.report_exporters import ReportDocument, ReportTable, to_csv, to_pdf, to_xlsx

GENERATE = "/api/v1/reports/generate"


@pytest.fixture(autouse=True)
def storage_dir(tmp_path: Path, monkeypatch) -> Path:
    """Every test writes report files into its own temporary directory."""
    directory = tmp_path / "reports"
    monkeypatch.setattr(settings, "REPORT_STORAGE_DIR", str(directory))
    return directory


def _seed(db_session: Session, patients: int = 12) -> None:
    for i in range(patients):
        patient = Patient(
            medical_record_number=f"MRN-RPT-{i}",
            age_group="61",
            gender="Female" if i % 2 else "Male",
            primary_diagnosis="Cardiac failure",
        )
        db_session.add(patient)
        db_session.commit()
        admission = Admission(
            patient_id=patient.id,
            department="Cardiology" if i % 3 else "Neurology",
            admission_date=date(2026, 1, 1 + i),
            discharge_date=date(2026, 1, 2 + i),
            time_in_hospital=3,
            readmitted="<30" if i % 4 == 0 else "NO",
            discharge_disposition="Home",
        )
        db_session.add(admission)
        db_session.commit()
        db_session.add(
            TreatmentOutcome(
                admission_id=admission.id,
                treatment_name="Insulin" if i % 2 else "Statins",
                outcome="improved" if i % 3 else "worsened",
                recovery_score=0.5 + i / 100,
            )
        )
        db_session.add(
            RiskPrediction(
                patient_id=patient.id,
                readmission_probability=0.8,
                risk_category="high" if i % 2 else "low",
                prediction_type="risk",
                model_name="risk",
                model_version="v1",
            )
        )
        db_session.commit()


def _generate(client: TestClient, header: dict[str, str], **body) -> dict:
    response = client.post(GENERATE, headers=header, json=body)
    assert response.status_code == 201, response.text
    return response.json()


def _stored_files(directory: Path) -> list[Path]:
    return sorted(directory.iterdir()) if directory.exists() else []


# --------------------------------------------------------------------------
# Generation, persistence and download - every type x every format
# --------------------------------------------------------------------------


@pytest.mark.parametrize("report_format", REPORT_FORMATS)
@pytest.mark.parametrize("report_type", REPORT_TYPES)
def test_every_report_type_generates_in_every_format(
    client: TestClient,
    auth_header,
    db_session: Session,
    storage_dir: Path,
    report_type: str,
    report_format: str,
) -> None:
    _seed(db_session)
    header = auth_header(Role.RESEARCHER)
    body = _generate(client, header, report_type=report_type, format=report_format)

    assert body["report_type"] == report_type
    assert body["format"] == report_format
    assert body["download_url"] == f"/api/v1/reports/{body['id']}/download"

    row = db_session.get(Report, body["id"])
    assert row is not None
    stored = report_storage.resolve(row.file_path)
    assert stored.parent == storage_dir.resolve()
    assert stored.stat().st_size == row.file_size_bytes == body["file_size_bytes"]

    download = client.get(body["download_url"], headers=header)
    assert download.status_code == 200
    assert download.content == stored.read_bytes()
    assert f'{report_type}_{body["id"]}_' in download.headers["content-disposition"]


def test_generated_report_is_persisted_with_its_filters_and_owner(
    client: TestClient, auth_header, db_session: Session
) -> None:
    _seed(db_session)
    body = _generate(
        client,
        auth_header(Role.HOSPITAL_ADMIN),
        report_type="treatment_effectiveness",
        format="csv",
        filters={"department": "Cardiology"},
    )
    row = db_session.get(Report, body["id"])
    assert row is not None
    assert row.filters == {"department": "Cardiology"}
    assert body["filters"] == {"department": "Cardiology"}
    assert row.generated_by == body["generated_by"]


def test_generation_is_audit_logged(client: TestClient, auth_header, db_session: Session) -> None:
    _seed(db_session)
    body = _generate(
        client, auth_header(Role.RESEARCHER), report_type="readmission_analytics", format="csv"
    )
    actions = db_session.execute(
        select(AuditLog.resource).where(AuditLog.action == "report.generate")
    ).scalars()
    assert f"report:{body['id']};type:readmission_analytics" in list(actions)


# --------------------------------------------------------------------------
# Export format content
# --------------------------------------------------------------------------


def test_csv_report_contains_service_figures(
    client: TestClient, auth_header, db_session: Session
) -> None:
    _seed(db_session)
    header = auth_header(Role.RESEARCHER)
    body = _generate(client, header, report_type="readmission_analytics", format="csv")
    text = client.get(body["download_url"], headers=header).content.decode("utf-8")
    rows = list(csv.reader(io.StringIO(text)))

    assert rows[0] == ["Report", "Readmission Analytics Report"]
    summary = {row[0]: row[1] for row in rows if len(row) == 2}
    assert summary["Patients"] == "12"
    assert summary["Admissions"] == "12"
    assert float(summary["Readmission rate"]) == pytest.approx(3 / 12)
    assert ["Month", "Admissions", "Readmission rate"] in rows


def _document(**overrides) -> ReportDocument:
    values = {
        "title": "Test Report",
        "generated_at": datetime(2026, 9, 27, 12, 0, tzinfo=UTC),
        "generated_by": "User #1 (researcher)",
        "filters": {},
        "summary": [("Patients", 12), ("Rate", 0.25)],
        "tables": [
            ReportTable(
                "Diagnoses",
                ["Diagnosis", "Count", "Share"],
                [['=HYPERLINK("x")', 3, 0.5], ["Diabète", 2, None], ["@SUM(A1)", 1, 0.1]],
            )
        ],
    }
    values.update(overrides)
    return ReportDocument(**values)


def test_csv_export_is_utf8_and_neutralises_formulas() -> None:
    text = to_csv(_document()).decode("utf-8")
    rows = list(csv.reader(io.StringIO(text)))
    data = {row[0]: row for row in rows if len(row) == 3}
    assert '\'=HYPERLINK("x")' in data
    assert "'@SUM(A1)" in data
    assert data["Diabète"] == ["Diabète", "2", ""]


def test_xlsx_export_has_sheets_headers_and_typed_values() -> None:
    workbook = load_workbook(io.BytesIO(to_xlsx(_document())))
    assert workbook.sheetnames == ["Summary", "Diagnoses"]

    sheet = workbook["Diagnoses"]
    assert [cell.value for cell in sheet[1]] == ["Diagnosis", "Count", "Share"]
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert isinstance(sheet["B2"].value, int)
    assert isinstance(sheet["C2"].value, float)
    # Stored as text, never as an executable formula.
    assert sheet["A2"].value == '\'=HYPERLINK("x")'
    assert sheet["A2"].data_type == "s"

    summary = workbook["Summary"]
    assert summary["A1"].value == "Test Report"
    values = {row[0]: row[1] for row in summary.iter_rows(values_only=True) if row[0]}
    assert values["Patients"] == 12


def test_xlsx_sheet_names_are_valid_and_unique() -> None:
    long_title = "Treatment/effectiveness: by [department]? " * 3
    document = _document(
        tables=[
            ReportTable(long_title, ["A"], [[1]]),
            ReportTable(long_title, ["A"], [[2]]),
        ]
    )
    names = load_workbook(io.BytesIO(to_xlsx(document))).sheetnames
    assert len(names) == len(set(names)) == 3
    assert all(len(name) <= 31 and not set(name) & set("[]:*?/\\") for name in names)


def test_pdf_export_is_a_titled_pdf_and_tolerates_non_latin_text() -> None:
    content = to_pdf(
        _document(
            title="Population Health Report",
            tables=[ReportTable("Diagnoses", ["Diagnosis"], [["糖尿病 (diabetes)"]])],
        )
    )
    assert content.startswith(b"%PDF-")
    assert b"/Title (Population Health Report)" in content
    assert b"/Type /Page" in content


def test_pdf_export_handles_an_empty_table() -> None:
    content = to_pdf(_document(tables=[ReportTable("Empty", ["A", "B"], [])]))
    assert content.startswith(b"%PDF-")


# --------------------------------------------------------------------------
# RBAC and ownership
# --------------------------------------------------------------------------


def test_doctors_cannot_use_reporting(client: TestClient, auth_header) -> None:
    """Documented gap: doctors hold no ANALYTICS_EXPORT."""
    header = auth_header(Role.DOCTOR)
    response = client.post(
        GENERATE, headers=header, json={"report_type": "treatment_effectiveness", "format": "csv"}
    )
    assert response.status_code == 403
    assert client.get("/api/v1/reports", headers=header).status_code == 403


@pytest.mark.parametrize("role", [Role.HOSPITAL_ADMIN, Role.RESEARCHER, Role.SYSTEM_ADMIN])
def test_export_roles_can_generate(
    client: TestClient, auth_header, db_session: Session, role: Role
) -> None:
    _seed(db_session)
    _generate(client, auth_header(role), report_type="population_health", format="pdf")


def test_reporting_requires_authentication(client: TestClient) -> None:
    assert client.get("/api/v1/reports").status_code == 401
    assert client.post(GENERATE, json={}).status_code == 401


def _without(monkeypatch, role: Role, permission: Permission) -> None:
    stripped = dict(rbac.PERMISSIONS)
    stripped[role] = rbac.PERMISSIONS[role] - {permission}
    monkeypatch.setattr(rbac, "PERMISSIONS", stripped)


def test_report_type_requires_its_analytics_permission(
    client: TestClient, auth_header, db_session: Session, monkeypatch, storage_dir: Path
) -> None:
    """ANALYTICS_EXPORT alone is not enough: the data permission is checked too."""
    _seed(db_session)
    _without(monkeypatch, Role.RESEARCHER, Permission.POPULATION_HEALTH_READ)
    response = client.post(
        GENERATE,
        headers=auth_header(Role.RESEARCHER),
        json={"report_type": "population_health", "format": "csv"},
    )
    assert response.status_code == 403
    assert _stored_files(storage_dir) == []


def test_download_rechecks_the_data_permission(
    client: TestClient, auth_header, db_session: Session, monkeypatch
) -> None:
    _seed(db_session)
    header = auth_header(Role.RESEARCHER)
    body = _generate(client, header, report_type="research_cohort", format="csv")
    _without(monkeypatch, Role.RESEARCHER, Permission.POPULATION_HEALTH_READ)
    assert client.get(body["download_url"], headers=header).status_code == 403


def test_other_users_reports_are_invisible(
    client: TestClient, auth_header, db_session: Session
) -> None:
    _seed(db_session)
    owner = auth_header(Role.RESEARCHER)
    other = auth_header(Role.HOSPITAL_ADMIN)
    body = _generate(client, owner, report_type="risk_distribution", format="csv")
    report_id = body["id"]

    assert client.get(f"/api/v1/reports/{report_id}", headers=other).status_code == 404
    assert client.get(f"/api/v1/reports/{report_id}/download", headers=other).status_code == 404
    assert client.delete(f"/api/v1/reports/{report_id}", headers=other).status_code == 404
    assert client.get("/api/v1/reports", headers=other).json() == []
    assert db_session.get(Report, report_id) is not None


def test_system_admin_sees_every_report(
    client: TestClient, auth_header, db_session: Session
) -> None:
    _seed(db_session)
    _generate(client, auth_header(Role.RESEARCHER), report_type="risk_distribution", format="csv")
    _generate(
        client, auth_header(Role.HOSPITAL_ADMIN), report_type="risk_distribution", format="csv"
    )
    response = client.get("/api/v1/reports", headers=auth_header(Role.SYSTEM_ADMIN))
    assert response.headers["X-Total-Count"] == "2"


def test_report_history_pages_and_filters_by_type(
    client: TestClient, auth_header, db_session: Session
) -> None:
    _seed(db_session)
    header = auth_header(Role.RESEARCHER)
    for report_format in REPORT_FORMATS:
        _generate(client, header, report_type="risk_distribution", format=report_format)
    _generate(client, header, report_type="readmission_analytics", format="csv")

    page = client.get("/api/v1/reports?limit=2", headers=header)
    assert page.headers["X-Total-Count"] == "4"
    assert len(page.json()) == 2

    filtered = client.get("/api/v1/reports?report_type=risk_distribution", headers=header)
    assert filtered.headers["X-Total-Count"] == "3"
    assert client.get("/api/v1/reports?report_type=nope", headers=header).status_code == 422


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "body",
    [
        {"report_type": "unknown", "format": "csv"},
        {"report_type": "population_health", "format": "docx"},
        {"report_type": "population_health"},
        {"report_type": "population_health", "format": "csv", "owner": 1},
        {"report_type": "population_health", "format": "csv", "filters": {"months": 6}},
        {"report_type": "risk_distribution", "format": "csv", "filters": {"months": 0}},
        {"report_type": "risk_distribution", "format": "csv", "filters": {"months": 61}},
        {"report_type": "risk_distribution", "format": "csv", "filters": {"colour": "red"}},
        {
            "report_type": "department_performance",
            "format": "csv",
            "filters": {"date_from": "2026-03-01", "date_to": "2026-01-01"},
        },
        {"report_type": "treatment_effectiveness", "format": "csv", "filters": {"department": ""}},
        {"report_type": "research_cohort", "format": "csv", "filters": {"diagnosis": "x" * 256}},
    ],
)
def test_invalid_generate_requests_are_rejected(
    client: TestClient, auth_header, body: dict, storage_dir: Path
) -> None:
    response = client.post(GENERATE, headers=auth_header(Role.SYSTEM_ADMIN), json=body)
    assert response.status_code == 422
    assert isinstance(response.json()["detail"], list)
    assert _stored_files(storage_dir) == []


def test_unsupported_filter_error_names_the_filter(client: TestClient, auth_header) -> None:
    response = client.post(
        GENERATE,
        headers=auth_header(Role.SYSTEM_ADMIN),
        json={"report_type": "population_health", "format": "csv", "filters": {"months": 6}},
    )
    assert "months not supported for report_type 'population_health'" in response.text


def test_department_report_applies_the_date_window(
    client: TestClient, auth_header, db_session: Session
) -> None:
    _seed(db_session)
    header = auth_header(Role.HOSPITAL_ADMIN)
    body = _generate(
        client,
        header,
        report_type="department_performance",
        format="csv",
        filters={"date_from": "2026-01-01", "date_to": "2026-01-03"},
    )
    rows = list(csv.reader(io.StringIO(client.get(body["download_url"], headers=header).text)))
    summary = {row[0]: row[1] for row in rows if len(row) == 2}
    assert summary["Admissions"] == "3"


# --------------------------------------------------------------------------
# Cohort guards carry through to reports
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("report_type", "filters"),
    [
        ("population_health", {}),
        ("research_cohort", {}),
        ("research_cohort", {"diagnosis": "Cardiac failure", "gender": "Male"}),
    ],
)
def test_small_cohorts_are_refused_and_leave_nothing_behind(
    client: TestClient,
    auth_header,
    db_session: Session,
    storage_dir: Path,
    report_type: str,
    filters: dict,
) -> None:
    _seed(db_session, patients=12 if filters else 3)
    response = client.post(
        GENERATE,
        headers=auth_header(Role.RESEARCHER),
        json={"report_type": report_type, "format": "xlsx", "filters": filters},
    )
    assert response.status_code == 422
    assert response.json()["detail"]["error"] == "cohort_too_small"
    assert db_session.execute(select(Report)).first() is None
    assert _stored_files(storage_dir) == []


def test_population_health_report_uses_generalised_age_bands(
    client: TestClient, auth_header, db_session: Session
) -> None:
    _seed(db_session)
    header = auth_header(Role.RESEARCHER)
    body = _generate(client, header, report_type="population_health", format="csv")
    text = client.get(body["download_url"], headers=header).text
    assert "60-69,12" in text
    assert "\n61," not in text


# --------------------------------------------------------------------------
# File handling
# --------------------------------------------------------------------------


def test_stored_file_names_are_unique_and_server_generated(
    client: TestClient, auth_header, db_session: Session
) -> None:
    _seed(db_session)
    header = auth_header(Role.RESEARCHER)
    first = _generate(client, header, report_type="risk_distribution", format="csv")
    second = _generate(client, header, report_type="risk_distribution", format="csv")
    names = {db_session.get(Report, r["id"]).file_path for r in (first, second)}  # type: ignore[union-attr]
    assert len(names) == 2
    assert all(len(name) == len("0" * 32 + ".csv") for name in names)


@pytest.mark.parametrize(
    "name",
    ["../secret.csv", "..\\secret.csv", "/etc/passwd", "report.csv", "a" * 32 + ".exe", ""],
)
def test_storage_refuses_names_it_did_not_generate(name: str) -> None:
    with pytest.raises(report_storage.InvalidReportFileNameError):
        report_storage.resolve(name)


def test_storage_never_overwrites_an_existing_file() -> None:
    name = report_storage.new_file_name("csv")
    report_storage.write(name, b"first")
    with pytest.raises(FileExistsError):
        report_storage.write(name, b"second")
    assert report_storage.resolve(name).read_bytes() == b"first"


def test_file_is_removed_when_the_metadata_insert_fails(
    db_session: Session, storage_dir: Path, monkeypatch
) -> None:
    from app.api.deps import CurrentUser
    from app.models.user import User
    from app.schemas.report import ReportGenerateRequest
    from app.services.report_service import ReportService

    _seed(db_session)
    user = User(
        email="r@hospital.org",
        full_name="R",
        hashed_password="x",
        role=str(Role.RESEARCHER),
        is_active=True,
    )
    db_session.add(user)
    db_session.commit()

    service = ReportService(db_session)

    def fail(**_: object) -> Report:
        raise RuntimeError("database unavailable")

    monkeypatch.setattr(service.reports, "create", fail)
    with pytest.raises(RuntimeError):
        service.generate(
            CurrentUser(subject=str(user.id), role=Role.RESEARCHER),
            ReportGenerateRequest(report_type="risk_distribution", format="pdf"),
        )
    assert _stored_files(storage_dir) == []


def test_download_of_a_missing_file_is_410(
    client: TestClient, auth_header, db_session: Session
) -> None:
    _seed(db_session)
    header = auth_header(Role.RESEARCHER)
    body = _generate(client, header, report_type="risk_distribution", format="csv")
    row = db_session.get(Report, body["id"])
    assert row is not None
    report_storage.resolve(row.file_path).unlink()
    assert client.get(body["download_url"], headers=header).status_code == 410


def test_a_tampered_file_path_is_never_served(
    client: TestClient, auth_header, db_session: Session
) -> None:
    _seed(db_session)
    header = auth_header(Role.RESEARCHER)
    body = _generate(client, header, report_type="risk_distribution", format="csv")
    row = db_session.get(Report, body["id"])
    assert row is not None
    row.file_path = "../../app/core/config.py"
    db_session.commit()

    assert client.get(body["download_url"], headers=header).status_code == 410
    assert client.delete(f"/api/v1/reports/{body['id']}", headers=header).status_code == 204


# --------------------------------------------------------------------------
# Deletion and retention purge
# --------------------------------------------------------------------------


def test_delete_removes_the_row_and_the_file(
    client: TestClient, auth_header, db_session: Session, storage_dir: Path
) -> None:
    _seed(db_session)
    header = auth_header(Role.HOSPITAL_ADMIN)
    body = _generate(client, header, report_type="department_performance", format="xlsx")
    assert len(_stored_files(storage_dir)) == 1

    assert client.delete(f"/api/v1/reports/{body['id']}", headers=header).status_code == 204
    assert _stored_files(storage_dir) == []
    assert client.get(f"/api/v1/reports/{body['id']}", headers=header).status_code == 404
    assert client.delete(f"/api/v1/reports/{body['id']}", headers=header).status_code == 404


def test_system_admin_can_delete_another_users_report(
    client: TestClient, auth_header, db_session: Session
) -> None:
    _seed(db_session)
    body = _generate(
        client, auth_header(Role.RESEARCHER), report_type="risk_distribution", format="csv"
    )
    response = client.delete(
        f"/api/v1/reports/{body['id']}", headers=auth_header(Role.SYSTEM_ADMIN)
    )
    assert response.status_code == 204


def test_purge_removes_only_expired_reports(
    client: TestClient, auth_header, db_session: Session, storage_dir: Path
) -> None:
    _seed(db_session)
    header = auth_header(Role.RESEARCHER)
    old = _generate(client, header, report_type="risk_distribution", format="csv")
    fresh = _generate(client, header, report_type="risk_distribution", format="pdf")
    row = db_session.get(Report, old["id"])
    assert row is not None
    row.generated_at = datetime.now(UTC) - timedelta(days=45)
    db_session.commit()

    response = client.post(
        "/api/v1/reports/purge?older_than_days=30", headers=auth_header(Role.SYSTEM_ADMIN)
    )
    assert response.status_code == 200
    assert response.json() == {"deleted": 1, "older_than_days": 30}
    db_session.expire_all()
    assert db_session.get(Report, old["id"]) is None
    assert db_session.get(Report, fresh["id"]) is not None
    assert len(_stored_files(storage_dir)) == 1


def test_purge_defaults_to_the_configured_retention(client: TestClient, auth_header) -> None:
    response = client.post("/api/v1/reports/purge", headers=auth_header(Role.SYSTEM_ADMIN))
    assert response.json() == {"deleted": 0, "older_than_days": settings.REPORT_RETENTION_DAYS}


@pytest.mark.parametrize("role", [Role.HOSPITAL_ADMIN, Role.RESEARCHER, Role.DOCTOR])
def test_purge_is_system_admin_only(client: TestClient, auth_header, role: Role) -> None:
    assert client.post("/api/v1/reports/purge", headers=auth_header(role)).status_code == 403


def test_purge_validates_the_window(client: TestClient, auth_header) -> None:
    response = client.post(
        "/api/v1/reports/purge?older_than_days=0", headers=auth_header(Role.SYSTEM_ADMIN)
    )
    assert response.status_code == 422
